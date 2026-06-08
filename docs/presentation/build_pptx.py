# -*- coding: utf-8 -*-
"""알비(Albi) 발표 PPTX 생성기 — python-pptx 기반.
브랜드 컬러 적용 + 슬라이드별 발표 대본을 notes 에 삽입.
실행: python docs/presentation/build_pptx.py
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---- 컬러 ----
PEAT   = RGBColor(0x15, 0x11, 0x0D)
CREAM  = RGBColor(0xFB, 0xF6, 0xEE)
AMBER  = RGBColor(0xB8, 0x73, 0x1A)
GOLD   = RGBColor(0xE0, 0xA8, 0x57)
DARK   = RGBColor(0x2A, 0x21, 0x18)
MUTED  = RGBColor(0x7A, 0x6A, 0x57)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
CARDBG = RGBColor(0xF3, 0xE9, 0xD8)

KFONT = "맑은 고딕"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = 13.333, 7.5


def add_slide(bg=CREAM):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = bg
    return s


def rect(slide, l, t, w, h, fill=None, line=None, line_w=1.0, shape=MSO_SHAPE.RECTANGLE):
    sp = slide.shapes.add_shape(shape, Inches(l), Inches(t), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(line_w)
    sp.shadow.inherit = False
    return sp


def text(slide, s, l, t, w, h, size=18, color=DARK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, font=KFONT, italic=False, spacing=1.0):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    lines = s.split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = spacing
        r = p.add_run(); r.text = ln
        r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
        r.font.name = font; r.font.color.rgb = color
    return tb


def bullets(slide, items, l, t, w, h, size=18, color=DARK, gap=10, marker="●  ", mcolor=AMBER):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap); p.line_spacing = 1.05
        rm = p.add_run(); rm.text = marker
        rm.font.size = Pt(size); rm.font.name = KFONT; rm.font.color.rgb = mcolor; rm.font.bold = True
        rt = p.add_run(); rt.text = it
        rt.font.size = Pt(size); rt.font.name = KFONT; rt.font.color.rgb = color
    return tb


def chrome(slide, title, message=None, page=None, accent=AMBER):
    """공통 상단 타이틀 + 좌측 액센트 바 + 푸터."""
    rect(slide, 0, 0, 0.22, SH, fill=accent)                       # 좌측 바
    text(slide, title, 0.85, 0.5, 11.5, 0.9, size=33, color=AMBER, bold=True)
    if message:
        text(slide, message, 0.88, 1.45, 11.8, 0.6, size=18, color=MUTED, italic=True)
    rect(slide, 0.9, 2.02, 4.2, 0.03, fill=GOLD)                   # 구분선
    # 푸터
    text(slide, "알비(Albi) · 모바일 앱 프로그래밍", 0.85, SH - 0.5, 7, 0.35,
         size=10, color=MUTED)
    if page is not None:
        text(slide, str(page), SW - 1.2, SH - 0.5, 0.6, 0.35, size=11, color=AMBER,
             bold=True, align=PP_ALIGN.RIGHT)


def notes(slide, txt):
    slide.notes_slide.notes_text_frame.text = txt


# =========================================================
# S1 — 표지
s = add_slide(PEAT)
rect(s, 0, 0, SW, 0.35, fill=AMBER)
rect(s, 0, SH - 0.35, SW, 0.35, fill=AMBER)
text(s, "알비 (Albi)", 1.0, 2.35, 11.3, 1.3, size=64, color=GOLD, bold=True, align=PP_ALIGN.CENTER)
text(s, "자연어로 한 줄 쓰면 끝나는 AI 음주 기록 비서", 1.0, 3.75, 11.3, 0.8,
     size=24, color=CREAM, align=PP_ALIGN.CENTER)
text(s, "모바일 앱 프로그래밍    |    팀  김태겸 · 곽지한    |    Flutter",
     1.0, 5.6, 11.3, 0.5, size=15, color=GOLD, align=PP_ALIGN.CENTER)
notes(s, "안녕하세요. 저희는 '알비'를 만든 김태겸, 곽지한입니다. 알비는 '알코올 비서'의 줄임말인데요, "
         "술 마신 걸 그냥 한 줄로 말하듯 적으면 AI가 알아서 정리해 주는 기록 앱입니다. Flutter로 만들었고요, "
         "지금부터 어떤 문제를 어떻게 풀었는지 보여드리겠습니다.")

# S2 — 문제 정의
s = add_slide()
chrome(s, "왜 만들었나", "음주 기록은 귀찮아서 작심삼일이 된다", page=2)
bullets(s, [
    "술 종류·양·도수·안주·장소를 매번 항목별로 입력 — 번거로움",
    "기록이 한 번 끊기면 그동안 쌓인 데이터가 무의미해짐",
    "정작 사용자가 궁금한 건 '내가 요즘 얼마나 마셨나' 하나뿐",
], 1.0, 2.6, 11.3, 3.5, size=21, gap=18)
notes(s, "음주 기록 앱들 한 번씩 써보셨을 텐데, 보통 작심삼일로 끝납니다. 마실 때마다 술 종류 고르고, 양, "
         "도수, 안주, 장소까지 일일이 입력해야 하거든요. 기록이 끊기면 데이터가 의미가 없어집니다. 정작 "
         "궁금한 건 '내가 요즘 얼마나 마셨지?' 하나인데, 그걸 보려면 매번 귀찮은 입력을 견뎌야 하는 "
         "모순이 있었습니다. 저희는 이 '입력 마찰'을 없애는 데 집중했습니다.")

# S3 — 해결 흐름
s = add_slide()
chrome(s, "해결: 한 줄이면 끝", "\"어제 강남에서 소주 2병에 삼겹살\"  →  자동으로 구조화", page=3)
steps = ["자연어 입력", "AI 파싱", "검토 · 수정", "저장"]
bw, gap, x0, y = 2.6, 0.55, 1.05, 3.4
for i, st in enumerate(steps):
    x = x0 + i * (bw + gap)
    box = rect(s, x, y, bw, 1.3, fill=CARDBG, line=AMBER, line_w=1.5, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    tf = box.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = st; r.font.size = Pt(19); r.font.bold = True
    r.font.name = KFONT; r.font.color.rgb = DARK
    if i < len(steps) - 1:
        text(s, "→", x + bw + 0.02, y + 0.35, gap, 0.6, size=28, color=AMBER, bold=True, align=PP_ALIGN.CENTER)
text(s, "입력의 번거로움은 없애되, 정확성은 '검토' 단계로 잡는다", 1.05, 5.3, 11, 0.6,
     size=17, color=MUTED, italic=True)
notes(s, "알비의 핵심은 이겁니다. '어제 강남에서 소주 2병에 삼겹살' — 이렇게 말하듯 한 줄만 쓰면 됩니다. "
         "그러면 앱이 소주, 2병, 도수, 안주는 삼겹살, 장소는 강남, 날짜는 어제로 쪼개서 카드로 보여줘요. "
         "흐름은 네 단계입니다. 자연어 입력, AI 파싱, 검토·수정, 저장. 입력의 번거로움은 없애되 정확성은 "
         "검토 단계로 잡는 구조입니다.")

# S4 — 차별점 3가지
s = add_slide()
chrome(s, "핵심 차별점 3가지", page=4)
cards = [
    ("①  자연어 AI 파싱", "폼이 아니라\n말하듯 입력"),
    ("②  AI 없어도 동작", "규칙 기반 파서가\n항상 fallback"),
    ("③  저장 전 검토", "자동 저장 금지\n안전 설계"),
]
cw, gap, x0, y = 3.5, 0.45, 1.05, 2.7
for i, (h, b) in enumerate(cards):
    x = x0 + i * (cw + gap)
    rect(s, x, y, cw, 2.9, fill=CARDBG, line=GOLD, line_w=1.5, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(s, x, y, cw, 0.12, fill=AMBER)
    text(s, h, x + 0.25, y + 0.45, cw - 0.5, 0.8, size=19, color=AMBER, bold=True)
    text(s, b, x + 0.25, y + 1.35, cw - 0.5, 1.3, size=17, color=DARK, spacing=1.1)
notes(s, "저희가 신경 쓴 세 가지입니다. 첫째, 자연어 AI 파싱 — 폼이 아니라 말로 입력해요. 둘째, 중요한데 "
         "AI가 없어도 동작합니다. 인터넷이 끊겼거나, API 키가 없거나, AI 호출이 실패해도 규칙 기반 파서가 "
         "곧바로 받아줍니다. 그래서 앱이 절대 먹통이 되지 않아요. 셋째, 저장 전에 반드시 검토 화면을 "
         "거칩니다. AI가 틀릴 수 있으니까요. AI를 쓰되 맹신하지 않는, 자동 저장을 일부러 막아둔 안전 설계입니다.")

# S5 — 아키텍처
s = add_slide()
chrome(s, "아키텍처 — MVVM + Riverpod", "역할이 분리된 단방향 구조", page=5)
layers = [
    ("View", "화면 표시만 담당"),
    ("ViewModel", "상태 · 로직 (Riverpod)"),
    ("Repository", "데이터베이스 접근"),
    ("SQLite", "로컬 저장"),
]
lh, gap, y0, x = 0.92, 0.28, 2.55, 1.4
for i, (nm, desc) in enumerate(layers):
    y = y0 + i * (lh + gap)
    shade = [AMBER, GOLD, RGBColor(0xD9, 0xBE, 0x8E), CARDBG][i]
    tcol = WHITE if i == 0 else DARK
    rect(s, x, y, 6.2, lh, fill=shade, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, nm, x + 0.35, y + 0.12, 2.6, lh, size=20, color=tcol, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, desc, x + 3.0, y + 0.12, 3.0, lh, size=15, color=tcol, anchor=MSO_ANCHOR.MIDDLE)
    if i < len(layers) - 1:
        text(s, "↓", x + 2.9, y + lh - 0.18, 0.5, 0.5, size=20, color=AMBER, bold=True, align=PP_ALIGN.CENTER)
text(s, "화면이 DB를 직접 건드리지 않도록 강제\n→ 유지보수 · 분업 용이",
     8.2, 3.4, 4.4, 2.0, size=18, color=MUTED, spacing=1.2)
notes(s, "구조는 MVVM 패턴입니다. 화면(View)은 보여주기만 하고, 상태와 로직은 ViewModel이, DB 접근은 "
         "Repository가 담당합니다. 화살표가 한 방향이죠. 화면이 DB를 직접 건드리는 일이 없도록 강제했습니다. "
         "상태관리는 Riverpod을 썼고요. 역할을 분리하니 한 부분을 고쳐도 다른 데가 안 깨지고, 두 명이 "
         "나눠 작업하기도 훨씬 수월했습니다.")

# S6 — AI 파싱 엔진
s = add_slide()
chrome(s, "AI 파싱 엔진 (ParseOrchestrator)", "AI를 쓰되, AI에 의존하지 않는다", page=6)
bullets(s, [
    "입력 → AI 사용 가능 여부 확인 → 가능하면 Gemini 호출",
    "실패하면 규칙 기반 파서로 자동 전환 — 앱이 멈추지 않음",
    "재시도 · 취소(CancelToken) · API 사용량 할당량 관리",
    "'AI를 실제로 시도했는지'를 사용자에게 투명하게 표시",
    "파싱 로직은 화면/ViewModel과 분리된 별도 모듈",
], 1.0, 2.55, 11.4, 3.6, size=20, gap=15)
notes(s, "핵심 원칙은 'AI를 쓰되 AI에 의존하지 않는다'입니다. 입력이 들어오면 AI를 쓸 수 있는지 확인하고, "
         "가능하면 Gemini를 호출합니다. 실패하면 그 순간 규칙 기반 파서로 자동 전환됩니다. 호출이 너무 "
         "오래 걸리면 취소할 수 있고, 사용량 할당량도 관리합니다. AI를 실제로 시도했는지 투명하게 알려줘서 "
         "결과를 믿을지 판단하게 했습니다. 이 파싱 로직은 화면이나 ViewModel이 아니라 별도 모듈로 "
         "완전히 분리해 뒀습니다.")

# S7 — 데이터 설계
s = add_slide()
chrome(s, "데이터 설계", "안정적인 로컬 우선 데이터베이스", page=7)
rect(s, 8.4, 2.5, 4.0, 3.0, fill=CARDBG, line=AMBER, line_w=1.5, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, "553", 8.4, 2.9, 4.0, 1.5, size=80, color=AMBER, bold=True, align=PP_ALIGN.CENTER)
text(s, "종 술 마스터 시드\n(한국 시장 브랜드 포함)", 8.4, 4.35, 4.0, 1.0, size=16,
     color=DARK, align=PP_ALIGN.CENTER, spacing=1.1)
bullets(s, [
    "SQLite 8개 테이블 + 트랜잭션으로 원자성 보장",
    "음식 사전 36종 · 장소 키워드 사전",
    "수정 시 기록 ID 보존 — 연결 데이터(테이스팅 노트) 유실 방지",
], 1.0, 2.7, 7.0, 3.2, size=19, gap=18)
notes(s, "데이터는 로컬 우선으로 SQLite를 썼습니다. 테이블이 8개인데, 기록 하나를 저장할 때 술, 안주, "
         "파싱 로그가 한꺼번에 들어가야 해서 트랜잭션으로 묶어 원자성을 보장했습니다. 술 데이터가 미리 "
         "553종 들어 있습니다. 위스키, 와인, 맥주부터 소주, 막걸리까지 한국 시장 브랜드를 직접 큐레이션해서 "
         "넣었어요. 음식 사전도 36종 있고요. 또 기록을 수정할 때 항목 ID를 보존해서, 연결된 테이스팅 노트 "
         "같은 데이터가 유실되지 않게 했습니다.")

# S8 — 화면 둘러보기
s = add_slide()
chrome(s, "주요 화면 둘러보기", "기록부터 회고까지", page=8)
screens = ["홈 (입력)", "검토 · 수정", "기록 목록", "캘린더", "아카이브 (술별)", "통계"]
cw, ch, gx, gy, x0, y0 = 3.6, 1.55, 0.35, 0.35, 1.05, 2.55
for i, nm in enumerate(screens):
    r, c = divmod(i, 3)
    x = x0 + c * (cw + gx); y = y0 + r * (ch + gy)
    rect(s, x, y, cw, ch, fill=CARDBG, line=GOLD, line_w=1.2, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, "[ 스크린샷 ]", x, y + 0.28, cw, 0.5, size=13, color=MUTED, align=PP_ALIGN.CENTER)
    text(s, nm, x, y + 0.82, cw, 0.5, size=17, color=AMBER, bold=True, align=PP_ALIGN.CENTER)
notes(s, "실제 화면입니다. 홈에서 한 줄 입력하고, 검토 화면에서 파싱 결과를 확인·수정합니다. 기록 목록은 "
         "월별로 묶이고, 캘린더로도 음주일을 한눈에 봅니다. 아카이브는 어떤 술을 얼마나 마셨나를 술 종류별로 "
         "모아주고, 통계 화면은 카테고리 비중을 차트로 보여줍니다. 기록하는 즐거움부터 돌아보는 재미까지 "
         "한 흐름으로 이어지게 설계했습니다. (※ 실제 스크린샷으로 교체)")

# S9 — 건강 인사이트
s = add_slide()
chrome(s, "건강 인사이트 & 리마인더", "단순 기록을 넘어 행동으로", page=9)
bullets(s, [
    "표준잔 계산 (한국 기준 1잔 = 순알코올 8g)으로 과음 신호 감지",
    "카테고리별 통계 차트로 음주 패턴 시각화",
    "로컬 알림 4종 — 주간 요약 · 3일 미기록 리마인드 ·",
    "    새벽 음주 다음날 입력 · 주간 과음 경고",
], 1.0, 2.6, 11.4, 3.4, size=20, gap=15)
notes(s, "알비는 기록을 넘어 행동까지 돕습니다. 마신 술을 한국 기준 표준잔으로 환산하는데, 1 표준잔이 "
         "순알코올 8그램이에요. 이걸로 과음 신호를 잡습니다. 통계는 차트로 보여주고요. 로컬 알림이 네 가지 "
         "있습니다. 매주 일요일 주간 요약, 3일 동안 기록이 없으면 리마인드, 새벽에 마신 다음 날 낮에 입력 "
         "알림, 한 주에 너무 많이 마셨으면 건강 경고. 기록을 하게 만들고, 더 나아가 습관을 돌아보게 만드는 "
         "장치들입니다.")

# S10 — 기술적 도전
s = add_slide()
chrome(s, "기술적 도전과 해결", "정상보다 비정상을 먼저 설계했다 — failure-path first", page=10)
bullets(s, [
    "통일된 에러 계층(sealed class) — 예외 처리 누락을 컴파일 단계에서 차단",
    "개인정보 보호 — 기록 삭제 시 AI 원문 데이터까지 같은 트랜잭션에서 정리",
    "클라우드 백업은 선택적 — 키 없으면 자동 비활성, 앱 본체는 정상 동작",
], 1.0, 2.6, 11.4, 3.4, size=20, gap=20)
notes(s, "가장 신경 쓴 원칙은 failure-path first, 정상 동작보다 비정상 상황을 먼저 설계하는 것이었습니다. "
         "첫째, 에러를 sealed class로 한 곳에 모아서 새 에러가 생기면 처리 누락이 컴파일 단계에서 잡히게 "
         "했습니다. 둘째, 개인정보 보호인데 기록을 삭제하면 AI에 보냈던 원문 데이터까지 같은 트랜잭션에서 "
         "정리됩니다. 셋째, 클라우드 백업은 선택 기능이라 키가 없으면 자동으로 꺼지고 앱 본체는 멀쩡히 "
         "동작합니다. 어떤 상황에서도 사용자 데이터가 안전한 쪽으로 설계했습니다.")

# S11 — 강의 개념 적용
s = add_slide()
chrome(s, "강의 개념의 적용", "수업에서 배운 핵심 주제를 실제로 구현", page=11)
pairs = [
    ("상태관리", "Riverpod"),
    ("로컬 DB", "sqflite vs Drift ORM 비교 (격리 데모 모듈)"),
    ("외부 REST API", "Gemini 연동 · 재시도 (Dio)"),
    ("권한 · 알림", "flutter_local_notifications"),
    ("클라우드", "Supabase (선택적 백업)"),
]
y = 2.6
for i, (k, v) in enumerate(pairs):
    rect(s, 1.05, y, 3.3, 0.62, fill=AMBER, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, k, 1.05, y, 3.3, 0.62, size=17, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, "→", 4.45, y, 0.5, 0.62, size=20, color=GOLD, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, v, 5.1, y, 7.3, 0.62, size=17, color=DARK, anchor=MSO_ANCHOR.MIDDLE)
    y += 0.78
notes(s, "이 프로젝트엔 수업에서 배운 개념이 거의 다 들어가 있습니다. 상태관리는 Riverpod으로, 로컬 DB는 "
         "sqflite로 구현하면서 같은 기능을 Drift라는 ORM으로도 따로 만들어 두 방식을 비교하는 격리 데모 "
         "모듈을 넣었습니다. 외부 REST API는 Gemini 호출에 재시도까지 붙여 다뤘고, 권한과 로컬 알림, "
         "클라우드 연동까지 — 강의 주제를 실제 제품 안에서 직접 적용해 본 게 가장 큰 수확이었습니다.")

# S12 — 팀 & 마무리
s = add_slide(PEAT)
rect(s, 0, 0, SW, 0.35, fill=AMBER)
text(s, "팀 역할 & 마무리", 0.9, 0.55, 11.5, 0.9, size=33, color=GOLD, bold=True)
# 두 컬럼
rect(s, 1.0, 1.9, 5.4, 1.9, fill=RGBColor(0x24,0x1D,0x15), line=AMBER, line_w=1.2, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, "김태겸", 1.3, 2.1, 4.8, 0.6, size=22, color=GOLD, bold=True)
text(s, "자연어 파싱 · 데이터베이스 · AI 연동", 1.3, 2.85, 4.8, 0.8, size=16, color=CREAM)
rect(s, 6.9, 1.9, 5.4, 1.9, fill=RGBColor(0x24,0x1D,0x15), line=AMBER, line_w=1.2, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, "곽지한", 7.2, 2.1, 4.8, 0.6, size=22, color=GOLD, bold=True)
text(s, "UI · 디자인 · 사용자 경험 (UX)", 7.2, 2.85, 4.8, 0.8, size=16, color=CREAM)
text(s, "\"편리함과 정확함은 '검토 단계' 하나로 양립한다\"", 1.0, 4.2, 11.3, 0.7,
     size=19, color=GOLD, italic=True, align=PP_ALIGN.CENTER)
text(s, "향후 계획 — 클라우드 동기화 본격화 · 디자인 리뉴얼", 1.0, 5.0, 11.3, 0.5,
     size=15, color=MUTED, align=PP_ALIGN.CENTER)
text(s, "감사합니다 — 알비(Albi)", 1.0, 5.9, 11.3, 0.9, size=30, color=GOLD,
     bold=True, align=PP_ALIGN.CENTER)
notes(s, "마지막으로 역할 분담입니다. 김태겸 학생이 자연어 파싱과 데이터베이스, AI 연동을 맡았고, 저 "
         "곽지한은 UI와 디자인, 사용자 경험을 담당했습니다. 저희가 배운 건 편리함과 정확함은 검토 단계 "
         "하나로 양립할 수 있다는 점이었어요. 앞으로는 클라우드 동기화를 본격화하고 디자인도 더 다듬을 "
         "계획입니다. 이상, 한 줄이면 끝나는 음주 기록 앱 알비 발표를 마치겠습니다. 감사합니다.")

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "albi_presentation.pptx")
prs.save(OUT)
print("OK saved:", OUT)
print("slides:", len(prs.slides._sldIdLst))
